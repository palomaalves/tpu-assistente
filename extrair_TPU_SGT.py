"""
Extrator de Tabelas Processuais Unificadas (TPU) - SGT/CNJ
Fonte: https://www.cnj.jus.br/sgt/sgt_ws.php

Gera arquivos JSON para uso no Assistente TPU (Artifact no Claude),
além dos arquivos Excel completos.

Dependências:
    pip install requests openpyxl

Uso:
    python extrair_TPU_SGT.py

Saída em Documentos:
    classes_sgt_cnj.json       <- faz upload deste no Artifact
    assuntos_sgt_cnj.json      <- faz upload deste no Artifact
    movimentos_sgt_cnj.json    <- faz upload deste no Artifact
    classes_sgt_cnj.xlsx       <- planilha completa (uso normal)
    assuntos_sgt_cnj.xlsx      <- planilha completa (uso normal)
    movimentos_sgt_cnj.xlsx    <- planilha completa (uso normal)
"""

import sys
import time
import json
import requests
import xml.etree.ElementTree as ET
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Garante saída UTF-8 no terminal Windows (evita UnicodeEncodeError com caracteres acentuados)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)

SOAP_URL  = "https://www.cnj.jus.br/sgt/sgt_ws.php"
NAMESPACE = "https://www.cnj.jus.br/sgt/sgt_ws.php"
DELAY     = 0.20
TIMEOUT   = 30
RETRIES   = 5  # tentativas por request; backoff exponencial: 2s, 4s, 8s, 16s

DOCS = Path.home() / "Documents"
DOCS.mkdir(exist_ok=True)

TABELAS = {
    "C": ("Classes",    DOCS / "classes_sgt_cnj"),
    "A": ("Assuntos",   DOCS / "assuntos_sgt_cnj"),
    "M": ("Movimentos", DOCS / "movimentos_sgt_cnj"),
}

TIPO_ITEM = "C"


def _envelope(action, corpo):
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<SOAP-ENV:Envelope '
        '  xmlns:SOAP-ENV="http://schemas.xmlsoap.org/soap/envelope/" '
        f' xmlns:ns1="{NAMESPACE}">'
        '<SOAP-ENV:Body>'
        f'<ns1:{action}>{corpo}</ns1:{action}>'
        '</SOAP-ENV:Body>'
        '</SOAP-ENV:Envelope>'
    )


def _chamar(action, corpo):
    headers = {
        "Content-Type": "text/xml; charset=utf-8",
        "SOAPAction":   f"{NAMESPACE}#{action}",
    }
    for tentativa in range(1, RETRIES + 1):
        try:
            resp = requests.post(
                SOAP_URL,
                data=_envelope(action, corpo).encode("utf-8"),
                headers=headers,
                timeout=TIMEOUT,
            )
            resp.raise_for_status()
            time.sleep(DELAY)
            # Retorna bytes para que ET.fromstring preserve a declaração de encoding do XML
            return resp.content
        except requests.RequestException as e:
            if tentativa < RETRIES:
                espera = 2 ** tentativa  # 2, 4, 8, 16 segundos
                print(f"  Tentativa {tentativa}/{RETRIES} falhou ({e}). Aguardando {espera}s...", flush=True)
                time.sleep(espera)
            else:
                print(f"  Erro após {RETRIES} tentativas em {action}: {e}")
                return None


def _texto(el):
    return (el.text or "").strip() if el is not None else ""


def get_filhos(seq_item):
    xml_resp = _chamar(
        "getArrayFilhosItemPublicoWS",
        f"<seqItem>{seq_item}</seqItem><tipoItem>{TIPO_ITEM}</tipoItem>",
    )
    if not xml_resp:
        return []
    try:
        root = ET.fromstring(xml_resp)
    except ET.ParseError as e:
        print(f"  XML inválido (seq={seq_item}): {e}")
        return []
    itens, vistos = [], set()
    for item in root.iter(f"{{{NAMESPACE}}}ArvoreGenerica"):
        seq  = _texto(item.find("seq_elemento"))
        nome = _texto(item.find("dsc_elemento"))
        if seq and nome and seq not in vistos:
            vistos.add(seq)
            itens.append({
                "seq":        seq,
                "nome":       nome,
                "ativo":      _texto(item.find("situacao")) or "S",
                "glossario":  _texto(item.find("glossario")) or "",
                "tem_filhos": _texto(item.find("temFilhos")) == "1",
            })
    return itens


def percorrer_arvore(seq_item, nivel, nomes, codigos, registros, total):
    for filho in get_filhos(seq_item):
        cod   = filho["seq"]
        nome  = filho["nome"]
        ativo = filho["ativo"]
        gloss = filho["glossario"]

        registros.append({
            "codigo":     cod,
            "nome":       nome,
            "nivel":      nivel,
            "hierarquia": " > ".join(nomes + [nome]),
            "codigo_pai": codigos[-1] if codigos else "0",
            "nome_pai":   nomes[-1]   if nomes   else "RAIZ",
            "ativo":      "Sim" if ativo.upper() in ("S", "SIM", "1", "TRUE", "A") else "Nao",
            "glossario":  (gloss or "")[:500],
        })
        total[0] += 1
        if total[0] % 50 == 0:
            print(f"  {total[0]} itens... nivel {nivel}: {nome[:50]}", flush=True)

        if filho["tem_filhos"]:
            try:
                seq_int = int(cod)
            except (ValueError, TypeError):
                continue
            percorrer_arvore(seq_int, nivel + 1,
                             nomes + [nome], codigos + [cod],
                             registros, total)


def salvar_json(registros, caminho_base):
    arquivo = Path(str(caminho_base) + ".json")
    with open(arquivo, "w", encoding="utf-8") as f:
        json.dump(registros, f, ensure_ascii=False, separators=(",", ":"))
    kb = arquivo.stat().st_size // 1024
    print(f"  JSON: {arquivo}  ({kb} KB)")


def salvar_excel(registros, caminho_base, nome_tabela):
    arquivo = Path(str(caminho_base) + ".xlsx")
    wb = Workbook()
    cab = ["Codigo", "Nome", "Nivel", "Hierarquia", "Cod_Pai", "Nome_Pai", "Ativo", "Glossario"]
    fill_h = PatternFill("solid", fgColor="1F4E79")
    font_h = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    borda  = Border(left=Side(style="thin"), right=Side(style="thin"),
                    top=Side(style="thin"),  bottom=Side(style="thin"))
    cores  = {1: "D6E4F0", 2: "EBF3FB", 3: "F5F9FC", 4: "FDFEFE"}
    fd     = Font(name="Arial", size=9)
    larg   = [10, 45, 8, 80, 12, 40, 8, 60]

    ws = wb.active
    ws.title = nome_tabela
    ws.append(cab)
    for i in range(1, len(cab) + 1):
        c = ws.cell(1, i)
        c.fill = fill_h
        c.font = font_h
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
    ws.row_dimensions[1].height = 30

    for i, r in enumerate(registros, 2):
        ws.append([r["codigo"], r["nome"], r["nivel"], r["hierarquia"],
                   r["codigo_pai"], r["nome_pai"], r["ativo"], r["glossario"]])
        fill = PatternFill("solid", fgColor=cores.get(r["nivel"], "FFFFFF"))
        for j in range(1, len(cab) + 1):
            c = ws.cell(i, j)
            c.fill = fill
            c.font = fd
            c.border = borda
            c.alignment = Alignment(vertical="center", wrap_text=(j == 8))
        ws.row_dimensions[i].height = 15

    for j, l in enumerate(larg, 1):
        ws.column_dimensions[get_column_letter(j)].width = l
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(registros) + 1}"
    wb.save(arquivo)
    print(f"  Excel: {arquivo}")


def main():
    print("=" * 60)
    print("Extrator TPU — SGT/CNJ  (JSON + Excel)")
    print("=" * 60)

    xml_v = _chamar("getDataUltimaVersao", "")
    if xml_v:
        try:
            for el in ET.fromstring(xml_v).iter("return"):
                print(f"Versao das tabelas: {_texto(el)}")
                break
        except ET.ParseError:
            pass

    extraidas = []
    for tipo, (nome, caminho) in TABELAS.items():
        global TIPO_ITEM
        TIPO_ITEM = tipo
        print(f"\n{'=' * 60}\nExtraindo: {nome} (tipo='{tipo}')\n{'=' * 60}", flush=True)

        registros, total = [], [0]
        percorrer_arvore(0, 1, [], [], registros, total)

        if not registros:
            print(f"  Nenhum registro encontrado — servidor indisponivel.")
            continue

        print(f"\n  Total {nome.lower()}: {len(registros)}")
        salvar_json(registros, caminho)
        salvar_excel(registros, caminho, nome)
        extraidas.append(nome)

    print(f"\n{'=' * 60}")
    if not extraidas:
        print("AVISO: Nenhuma tabela foi extraida. Servidor CNJ indisponivel.")
    else:
        print(f"Extracao concluida! Tabelas salvas: {', '.join(extraidas)}")
        print(f"Arquivos em: {DOCS}")
        print("\nProximo passo — no Assistente TPU (Artifact):")
        print("  Aba 'Carregar Tabelas' > Upload manual")
        print("  Selecione os 3 arquivos .json gerados")
    print("=" * 60)


if __name__ == "__main__":
    main()
